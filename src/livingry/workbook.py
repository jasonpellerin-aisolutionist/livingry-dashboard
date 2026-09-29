"""Excel workbook a reader can open without Tableau or Python.

The ratio cells are formulas, so the grouping is visible in the sheet rather
than hidden in a script. Navy and teal only.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from livingry import EXPORTS, REPORTS

NAVY = "1E3A5F"
TEAL = "0F766E"
SLATE = "64748B"
INK = "0F172A"
WHITE = "FFFFFF"
BAND = "F1F5F9"

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color=WHITE)
TITLE_FONT = Font(bold=True, size=16, color=INK)
NOTE_FONT = Font(italic=True, size=10, color=SLATE)


def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(labels))}{row}"


def _widths(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def main() -> None:
    agencies = pd.read_csv(EXPORTS / "agency_year.csv")
    claims = pd.read_csv(EXPORTS / "claims.csv")
    series = pd.read_csv(EXPORTS / "series.csv")
    protein = pd.read_csv(EXPORTS / "protein_land.csv")
    gaps = pd.read_csv(EXPORTS / "not_measured.csv")
    checks = pd.read_csv(EXPORTS / "quality_checks.csv")
    co2 = pd.read_csv(EXPORTS / "co2_monthly.csv")

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = "The Livingry Dashboard"
    readme["A1"].font = TITLE_FONT
    readme["A2"] = (
        "Indicators named in The Livingry Protocol, an unpublished manuscript by Jason Pellerin. "
        "Each number is in its own unit. This workbook does not compute a score."
    )
    readme["A2"].alignment = Alignment(wrap_text=True)
    readme.merge_cells("A2:B2")
    notes = [
        "Obligations are commitments, not cash. FY2026 is a partial year and is not used in the claims.",
        "The USDA + EPA + HUD + Interior total is an author grouping. The formula is in the Priority sheet.",
        "HHS is shown and never added in. Most of it is Medicare and Medicaid.",
        "A Lockheed registration and the EPA total are different kinds of quantity. Scale only.",
        "Calories beside food insecurity is coexistence, not a cause.",
        "GDP beside CO2 is two series, not an index.",
        "Land per 100g protein is Poore and Nemecek 2018 via Our World in Data, not a live feed.",
        "Chapter 1 refuses a vacant-homes versus homeless-people comparison. It is not here.",
        "Chapter 29 refuses a single composite. It is not here.",
    ]
    for i, note in enumerate(notes, start=4):
        readme.cell(row=i, column=1, value=note).font = NOTE_FONT
    _widths(readme, [110])
    readme.row_dimensions[2].height = 36

    ws = wb.create_sheet("Claims")
    _header(ws, 1, ["Chapter", "Claim", "Verdict", "Year", "Evidence", "Caveat"])
    for i, row in enumerate(claims.itertuples(index=False), start=2):
        values = [row.book_chapter, row.statement, row.verdict, row.year, row.evidence, row.caveat]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 48
    _widths(ws, [14, 42, 24, 10, 55, 45])

    ws = wb.create_sheet("Priority")
    complete = agencies[agencies["fiscal_year"] == 2025].sort_values("obligations", ascending=False)
    _header(ws, 1, ["Agency", "Group", "FY2025 obligations", "Share of these eight"])
    for i, row in enumerate(complete.itertuples(index=False), start=2):
        ws.cell(row=i, column=1, value=row.agency)
        ws.cell(row=i, column=2, value=row.group)
        cell = ws.cell(row=i, column=3, value=row.obligations)
        cell.number_format = '"$"#,##0'
        ws.cell(row=i, column=4, value=f"=C{i}/SUM($C$2:$C$9)")
        ws.cell(row=i, column=4).number_format = "0.0%"
    last = 1 + len(complete)
    group_row = last + 2
    ws.cell(row=group_row, column=1, value="Author group: USDA, EPA, HUD, Interior").font = Font(bold=True, color=TEAL)
    ws.cell(row=group_row, column=3, value=(
        f'=SUMIF(A2:A{last},"Department of Agriculture",C2:C{last})'
        f'+SUMIF(A2:A{last},"Environmental Protection Agency",C2:C{last})'
        f'+SUMIF(A2:A{last},"Department of Housing and Urban Development",C2:C{last})'
        f'+SUMIF(A2:A{last},"Department of the Interior",C2:C{last})'
    ))
    ws.cell(row=group_row, column=3).number_format = '"$"#,##0'
    ws.cell(row=group_row + 1, column=1, value="Defense divided by that group").font = Font(bold=True, color=NAVY)
    ws.cell(row=group_row + 1, column=3, value=f'=SUMIF(A2:A{last},"Department of Defense",C2:C{last})/C{group_row}')
    ws.cell(row=group_row + 1, column=3).number_format = '0.0" to 1"'
    ws.cell(row=group_row + 3, column=1, value="HHS is on the list above and is not in the group. Most of it is mandatory.").font = NOTE_FONT
    _widths(ws, [62, 16, 28, 24])

    chart = BarChart()
    chart.type = "col"
    chart.title = "FY2025 obligations"
    chart.y_axis.title = "Dollars"
    data = Reference(ws, min_col=3, min_row=1, max_row=last)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.style = 10
    chart.series[0].graphicalProperties.solidFill = NAVY
    chart.width = 20
    chart.height = 8
    ws.add_chart(chart, "A16")

    ws = wb.create_sheet("Provision")
    keep = series[series["family"].isin(["provision"])].copy()
    latest = keep.sort_values("year").groupby(["entity", "indicator"], as_index=False).tail(1)
    _header(ws, 1, ["Place", "Indicator", "Year", "Value", "Unit", "Source"])
    for i, row in enumerate(latest.sort_values(["indicator", "entity"]).itertuples(index=False), start=2):
        for col, value in enumerate([row.entity, row.indicator, row.year, row.value, row.unit, row.source], start=1):
            ws.cell(row=i, column=col, value=value)
        ws.cell(row=i, column=4).number_format = "#,##0.00"
    _widths(ws, [20, 42, 10, 16, 32, 55])

    ws = wb.create_sheet("GDP beside CO2")
    annual = co2.groupby("year", as_index=False)["ppm"].mean()
    gdp = series[(series["entity"] == "United States") & (series["indicator"] == "gdp_per_capita_constant_2015_usd")]
    joined = gdp.merge(annual, on="year", how="inner").sort_values("year")
    _header(ws, 1, ["Year", "US GDP per capita, constant 2015 $", "Mauna Loa annual mean, ppm"])
    for i, row in enumerate(joined.itertuples(index=False), start=2):
        ws.cell(row=i, column=1, value=int(row.year))
        ws.cell(row=i, column=2, value=row.value).number_format = '"$"#,##0'
        ws.cell(row=i, column=3, value=row.ppm).number_format = "0.00"
    _widths(ws, [12, 40, 36])
    ws.cell(row=2, column=5, value="Two columns. No ratio, no index.").font = NOTE_FONT

    ws = wb.create_sheet("Protein land")
    _header(ws, 1, ["Food", "Year", "m2 per 100g protein", "Source"])
    for i, row in enumerate(protein.itertuples(index=False), start=2):
        ws.cell(row=i, column=1, value=row.entity)
        ws.cell(row=i, column=2, value=int(row.year) if pd.notna(row.year) else None)
        ws.cell(row=i, column=3, value=row.land_use_m2).number_format = "#,##0.00"
        ws.cell(row=i, column=4, value=row.source)
    _widths(ws, [28, 10, 24, 55])

    ws = wb.create_sheet("Not measured")
    _header(ws, 1, ["Indicator", "Chapter", "Why it is absent"])
    for i, row in enumerate(gaps.itertuples(index=False), start=2):
        for col, value in enumerate([row.indicator, row.book_chapter, row.why_it_is_absent], start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 32
    _widths(ws, [32, 14, 90])

    ws = wb.create_sheet("Quality log")
    _header(ws, 1, ["Check", "Passed", "Detail"])
    for i, row in enumerate(checks.itertuples(index=False), start=2):
        ws.cell(row=i, column=1, value=row.check_name)
        ws.cell(row=i, column=2, value="yes" if row.passed else "no")
        ws.cell(row=i, column=3, value=row.detail)
    _widths(ws, [52, 12, 24])

    out = REPORTS / "livingry_dashboard.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out.relative_to(REPORTS.parent)} ({out.stat().st_size / 1024:.0f} KB, {len(wb.sheetnames)} tabs)")


if __name__ == "__main__":
    main()
